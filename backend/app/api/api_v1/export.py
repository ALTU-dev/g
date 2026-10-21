import csv
import json
import re
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path
from typing import Literal
from urllib.parse import quote

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel

from app.core.security import get_current_user
from app.models.user import User

router = APIRouter()

DOCX_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "templates"
    / "altrag_document_template.docx"
)
OFFICIAL_TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "templates" / "official"
OFFICIAL_DOCX_TEMPLATES = {
    "ri": OFFICIAL_TEMPLATE_DIR / "ri_alt_125.docx",
    "pp": OFFICIAL_TEMPLATE_DIR / "form_pp_alt_125.docx",
    "di": OFFICIAL_TEMPLATE_DIR / "form_di_alt_125.docx",
}

GENERIC_TITLES = {"ALTBot answer", "altRAG answer", "answer"}
DEFAULT_ORGANIZATION = "АО «ALT Университет имени Мухамеджана Тынышпаева»"
DOCX_META_RE = re.compile(
    r"<!--\s*DOCX_META\s*(?P<body>.*?)\s*-->",
    flags=re.IGNORECASE | re.DOTALL,
)

META_KEY_ALIASES = {
    "template": "template",
    "document_template": "template",
    "doc_template": "template",
    "шаблон": "template",
    "форма": "template",
    "document_type": "document_type",
    "doc_type": "document_type",
    "type": "document_type",
    "тип документа": "document_type",
    "құжат түрі": "document_type",
    "title": "title",
    "название": "title",
    "заголовок": "title",
    "тақырып": "title",
    "organization": "organization",
    "организация": "organization",
    "ұйым": "organization",
    "approval_required": "approval_required",
    "requires_approval": "approval_required",
    "утверждение": "approval_required",
    "требуется утверждение": "approval_required",
    "approval_authority": "approval_authority",
    "approval_type": "approval_authority",
    "кем утверждается": "approval_authority",
    "орган утверждения": "approval_authority",
    "board_protocol": "board_protocol",
    "protocol": "board_protocol",
    "протокол": "board_protocol",
    "номер протокола": "board_protocol",
    "approver_position": "approver_position",
    "approval_position": "approver_position",
    "должность утверждающего": "approver_position",
    "должность": "approver_position",
    "approver_name": "approver_name",
    "approval_name": "approver_name",
    "фио утверждающего": "approver_name",
    "фио": "approver_name",
    "recipient_position": "recipient_position",
    "должность адресата": "recipient_position",
    "recipient_name": "recipient_name",
    "адресат": "recipient_name",
    "кому": "recipient_name",
    "фио адресата": "recipient_name",
    "author_position": "author_position",
    "должность автора": "author_position",
    "author_name": "author_name",
    "автор": "author_name",
    "от кого": "author_name",
    "фио автора": "author_name",
    "date": "date",
    "дата": "date",
    "күні": "date",
    "city_year": "city_year",
    "город_год": "city_year",
    "document_code": "document_code",
    "code": "document_code",
    "код документа": "document_code",
    "номер документа": "document_code",
    "department_full_name": "department_full_name",
    "full_department_name": "department_full_name",
    "полное наименование подразделения": "department_full_name",
    "department_short_name": "department_short_name",
    "short_department_name": "department_short_name",
    "сокращенное наименование подразделения": "department_short_name",
    "position_full_name": "position_full_name",
    "full_position_name": "position_full_name",
    "полное наименование должности": "position_full_name",
    "position_short_name": "position_short_name",
    "short_position_name": "position_short_name",
    "сокращенное наименование должности": "position_short_name",
    "supervisor_position": "supervisor_position",
    "curator_position": "supervisor_position",
    "должность руководителя": "supervisor_position",
    "должность руководителя/куратора": "supervisor_position",
    "leader_position": "leader_position",
    "должность руководителя подразделения": "leader_position",
    "extra_documents": "extra_documents",
    "прочие документы": "extra_documents",
    "purpose": "purpose",
    "цель": "purpose",
    "tasks": "tasks",
    "задачи": "tasks",
    "functions": "functions",
    "функции": "functions",
    "duties": "duties",
    "обязанности": "duties",
    "rights": "rights",
    "права": "rights",
    "qualification_requirements": "qualification_requirements",
    "квалификационные требования": "qualification_requirements",
}

FORMAL_DOCUMENT_KEYWORDS = (
    "приказ",
    "бұйрық",
    "письмо",
    "хат",
    "заявление",
    "өтініш",
    "служебная записка",
    "докладная",
    "положение",
    "ереже",
    "инструкция",
    "нұсқаулық",
    "регламент",
    "правила",
    "қағида",
    "договор",
    "шарт",
    "акт",
    "справка",
    "анықтама",
    "протокол",
    "хаттама",
    "отчет",
    "отчёт",
    "есеп",
)

APPROVAL_DOCUMENT_KEYWORDS = (
    "приказ",
    "бұйрық",
    "положение",
    "ереже",
    "инструкция",
    "нұсқаулық",
    "регламент",
    "правила",
    "қағида",
    "политика",
    "стандарт",
)

LETTER_DOCUMENT_KEYWORDS = (
    "письмо",
    "хат",
    "заявление",
    "өтініш",
    "служебная записка",
    "докладная",
)

EMPTY_META_VALUES = {"", "...", "…", "-", "—", "n/a", "нет", "не указано", "not provided"}

TEMPLATE_KEY_ALIASES = {
    "pp": "pp",
    "пп": "pp",
    "form_pp": "pp",
    "форма 1": "pp",
    "форма №1": "pp",
    "положение": "pp",
    "положение о подразделении": "pp",
    "di": "di",
    "ди": "di",
    "form_di": "di",
    "форма 2": "di",
    "форма №2": "di",
    "должностная инструкция": "di",
    "ri": "ri",
    "ри": "ri",
    "рабочая инструкция": "ri",
    "ри-alt-125": "ri",
}

TEMPLATE_SECTION_HEADINGS = {
    "pp": (
        "ОБЩИЕ ПОЛОЖЕНИЯ",
        "СТРУКТУРА",
        "ЗАДАЧИ И ФУНКЦИИ",
        "ОТВЕТСТВЕННОСТЬ",
        "ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ",
    ),
    "di": (
        "ОБЩИЕ ПОЛОЖЕНИЯ",
        "КВАЛИФИКАЦИОННЫЕ ТРЕБОВАНИЯ",
        "ОБЯЗАННОСТИ",
        "ПРАВА",
        "ОТВЕТСТВЕННОСТЬ",
        "ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ",
    ),
    "ri": (
        "ОБЩИЕ ПОЛОЖЕНИЯ",
        "РАЗРАБОТКА, СОГЛАСОВАНИЕ И УТВЕРЖДЕНИЕ ПОЛОЖЕНИЯ О ПОДРАЗДЕЛЕНИИ",
        "РАЗРАБОТКА, СОГЛАСОВАНИЕ И УТВЕРЖДЕНИЕ ДОЛЖНОСТНОЙ ИНСТРУКЦИИ",
        "ХРАНЕНИЕ, ВНЕСЕНИЕ ИЗМЕНЕНИЙ И ВВЕДЕНИЕ В ДЕЙСТВИЕ",
        "ОТВЕТСТВЕННОСТЬ",
        "ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ",
    ),
}

TEMPLATE_STOP_HEADINGS = {
    "ЛИСТ ОЗНАКОМЛЕНИЯ",
    "ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ",
}


class ExportRequest(BaseModel):
    content: str
    title: str = "ALTBot answer"


ExportFormat = Literal["docx", "xlsx", "txt", "md", "csv"]


def _safe_filename(title: str, extension: str) -> str:
    name = re.sub(r"[^\w\s.-]+", "", title, flags=re.UNICODE).strip()
    name = re.sub(r"\s+", " ", name)[:90].strip(" .")
    if not name:
        name = "ALTBot-answer"
    if not name.lower().endswith(f".{extension}"):
        name = f"{name}.{extension}"
    return name


def _ascii_fallback(filename: str) -> str:
    return "".join(
        char if 32 <= ord(char) < 127 and char not in {'"', "\\"} else "_"
        for char in filename
    ) or "ALTBot-answer"


def _download_response(
    data: bytes,
    filename: str,
    media_type: str,
) -> StreamingResponse:
    return StreamingResponse(
        BytesIO(data),
        media_type=media_type,
        headers={
            "Content-Disposition": (
                f'attachment; filename="{_ascii_fallback(filename)}"; '
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )


def _strip_markdown(text: str) -> str:
    cleaned = text
    cleaned = re.sub(r"```[\s\S]*?```", lambda match: match.group(0).strip("`"), cleaned)
    cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
    cleaned = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", cleaned)
    cleaned = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", cleaned)
    cleaned = re.sub(r"^\s{0,3}#{1,6}\s*", "", cleaned, flags=re.MULTILINE)
    cleaned = re.sub(r"^\s{0,3}>\s?", "", cleaned, flags=re.MULTILINE)
    cleaned = re.sub(r"(\*\*|__)(.*?)\1", r"\2", cleaned)
    cleaned = re.sub(r"(\*|_)(.*?)\1", r"\2", cleaned)
    return cleaned.strip()


def _set_run_font(run, *, bold: bool = False, italic: bool = False, size: int = 14) -> None:
    run.font.name = "Times New Roman"
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _normalize_meta_key(key: str) -> str:
    normalized = re.sub(r"\s+", " ", key.strip().lower().replace("-", "_"))
    return META_KEY_ALIASES.get(normalized, normalized)


def _parse_docx_meta_block(block: str) -> dict[str, str]:
    meta: dict[str, str] = {}
    stripped = block.strip()
    if not stripped:
        return meta

    if stripped.startswith("{"):
        try:
            parsed = json.loads(stripped)
            if isinstance(parsed, dict):
                return {
                    _normalize_meta_key(str(key)): str(value).strip()
                    for key, value in parsed.items()
                    if value is not None and str(value).strip().lower() not in EMPTY_META_VALUES
                }
        except json.JSONDecodeError:
            pass

    for raw_line in stripped.splitlines():
        line = raw_line.strip().strip("#").strip()
        if not line or line.startswith("//"):
            continue
        if ":" not in line and "=" not in line:
            continue
        separator = ":" if ":" in line else "="
        key, value = line.split(separator, 1)
        key = _normalize_meta_key(key)
        value = value.strip().strip('"').strip("'")
        if key and value.lower() not in EMPTY_META_VALUES:
            meta[key] = value
    return meta


def _extract_visible_docx_meta(content: str) -> tuple[dict[str, str], str]:
    meta: dict[str, str] = {}
    lines = content.splitlines()
    cleaned_lines: list[str] = []
    consumed_any = False
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if index > 24:
            break

        if not stripped and consumed_any:
            index += 1
            continue

        match = re.match(r"^([^:：]{1,48})[:：]\s*(.+)$", stripped)
        if not match:
            break

        key = _normalize_meta_key(match.group(1))
        if key not in META_KEY_ALIASES.values():
            break

        value = match.group(2).strip()
        if value and value.lower() not in EMPTY_META_VALUES:
            meta[key] = value
            consumed_any = True
        index += 1

    cleaned_lines.extend(lines[index:])
    return meta, "\n".join(cleaned_lines).strip()


def _extract_docx_meta(content: str) -> tuple[dict[str, str], str]:
    meta: dict[str, str] = {}

    for match in DOCX_META_RE.finditer(content):
        meta.update(_parse_docx_meta_block(match.group("body")))

    content_without_hidden_meta = DOCX_META_RE.sub("", content).strip()
    visible_meta, cleaned = _extract_visible_docx_meta(content_without_hidden_meta)
    visible_meta.update(meta)
    return visible_meta, cleaned


def _meta_value(meta: dict[str, str], *keys: str, default: str = "") -> str:
    for key in keys:
        value = meta.get(key, "").strip()
        if value and value.lower() not in EMPTY_META_VALUES:
            return value
    return default


def _coerce_meta_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    normalized = value.strip().lower()
    if normalized in {"true", "yes", "1", "да", "иә", "yes"}:
        return True
    if normalized in {"false", "no", "0", "нет", "жоқ"}:
        return False
    return None


def _first_markdown_heading(content: str) -> str:
    for raw_line in content.splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+)$", raw_line.strip())
        if match:
            return _strip_markdown(match.group(1))
    return ""


def _document_type_from(title: str, content: str, meta: dict[str, str]) -> str:
    return (
        _meta_value(meta, "document_type")
        or _first_markdown_heading(content)
        or ("" if title in GENERIC_TITLES else title)
        or "Документ"
    ).strip()


def _document_title_from(title: str, content: str, meta: dict[str, str], document_type: str) -> str:
    return (
        _meta_value(meta, "title")
        or ("" if title in GENERIC_TITLES else title)
        or _first_markdown_heading(content)
        or document_type
    ).strip()


def _contains_any(text: str, keywords: tuple[str, ...]) -> bool:
    lowered = text.lower()
    return any(keyword in lowered for keyword in keywords)


def _should_use_formal_docx(meta: dict[str, str], content: str, title: str) -> bool:
    if meta:
        return True
    return _contains_any(f"{title}\n{content[:1500]}", FORMAL_DOCUMENT_KEYWORDS)


def _approval_required(
    meta: dict[str, str],
    document_type: str,
    title: str,
    content: str,
) -> bool:
    explicit = _coerce_meta_bool(meta.get("approval_required"))
    if explicit is not None:
        return explicit

    approver_position = _meta_value(meta, "approver_position")
    approver_name = _meta_value(meta, "approver_name")
    approver_text = f"{approver_position} {approver_name}".lower()
    is_rector = "ректор" in approver_text
    is_approval_document = _contains_any(
        f"{document_type}\n{title}\n{content[:500]}",
        APPROVAL_DOCUMENT_KEYWORDS,
    )
    return bool(is_approval_document and is_rector and approver_name)


def _is_letter_like_document(document_type: str, title: str) -> bool:
    return _contains_any(f"{document_type}\n{title}", LETTER_DOCUMENT_KEYWORDS)


def _is_placeholder(value: str) -> bool:
    cleaned = value.strip()
    return (
        not cleaned
        or cleaned.lower() in EMPTY_META_VALUES
        or cleaned.startswith("[")
        or cleaned.startswith("*")
    )


def _placeholder(label: str) -> str:
    return f"[{label}]"


def _add_formal_run(paragraph, text: str, *, bold: bool = False, size: int = 14):
    run = paragraph.add_run(text)
    _set_run_font(run, bold=bold, size=size)
    if _is_placeholder(text):
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _add_formal_paragraph(
    document,
    text: str = "",
    *,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    bold: bool = False,
    size: int = 14,
    before: int = 0,
    after: int = 0,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    if text:
        _add_formal_run(paragraph, text, bold=bold, size=size)
    return paragraph


def _apply_formal_page_setup(document) -> None:
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)


def _add_spacer(document, points: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(points)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def _add_approval_block(document, meta: dict[str, str]) -> None:
    organization = _meta_value(meta, "organization", default=DEFAULT_ORGANIZATION)
    position = _meta_value(meta, "approver_position", default=_placeholder("должность"))
    name = _meta_value(meta, "approver_name", default=_placeholder("Ф.И.О."))
    doc_date = _meta_value(meta, "date", default=f"«___» __________ {date.today().year} г.")

    _add_formal_paragraph(
        document,
        "УТВЕРЖДАЮ",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        after=2,
    )
    _add_formal_paragraph(
        document,
        position,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        after=0,
    )
    _add_formal_paragraph(
        document,
        organization,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        after=0,
    )
    signature = _add_formal_paragraph(
        document,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        after=0,
    )
    _add_formal_run(signature, "__________________ ", bold=True)
    _add_formal_run(signature, name, bold=True)
    _add_formal_paragraph(
        document,
        doc_date,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        after=0,
    )


def _add_address_block(document, meta: dict[str, str]) -> None:
    recipient_position = _meta_value(meta, "recipient_position")
    recipient_name = _meta_value(meta, "recipient_name")
    author_position = _meta_value(meta, "author_position")
    author_name = _meta_value(meta, "author_name")

    if not any((recipient_position, recipient_name, author_position, author_name)):
        return

    if recipient_position or recipient_name:
        _add_formal_paragraph(
            document,
            "Кому:",
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            bold=True,
            after=0,
        )
        if recipient_position:
            _add_formal_paragraph(document, recipient_position, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if recipient_name:
            _add_formal_paragraph(document, recipient_name, align=WD_ALIGN_PARAGRAPH.RIGHT)

    if author_position or author_name:
        _add_formal_paragraph(
            document,
            "От:",
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            bold=True,
            before=6,
            after=0,
        )
        if author_position:
            _add_formal_paragraph(document, author_position, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if author_name:
            _add_formal_paragraph(document, author_name, align=WD_ALIGN_PARAGRAPH.RIGHT)


def _same_heading(left: str, right: str) -> bool:
    return _normalized_heading(left) == _normalized_heading(right)


def _remove_leading_duplicate_title(content: str, document_type: str, title: str) -> str:
    lines = content.splitlines()
    index = 0
    while index < len(lines) and not lines[index].strip():
        index += 1

    while index < len(lines):
        stripped = lines[index].strip()
        heading = re.match(r"^#{1,6}\s+(.+)$", stripped)
        candidate = heading.group(1).strip() if heading else stripped
        if candidate and (
            _same_heading(candidate, document_type)
            or _same_heading(candidate, title)
        ):
            index += 1
            while index < len(lines) and not lines[index].strip():
                index += 1
            continue
        break

    return "\n".join(lines[index:]).strip()


def _add_formal_title_block(
    document,
    document_type: str,
    title: str,
    *,
    standalone_title_page: bool,
) -> None:
    _add_spacer(document, 140 if standalone_title_page else 72)

    type_text = _strip_markdown(document_type).upper()
    _add_formal_paragraph(
        document,
        type_text,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=14,
        after=8,
    )

    if title and not _same_heading(title, document_type):
        _add_formal_paragraph(
            document,
            _strip_markdown(title),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            size=14,
            after=18,
        )


def _append_markdown_to_docx(document, content: str) -> None:
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.rstrip()
        if not line:
            document.add_paragraph()
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            _add_heading(document, heading.group(2).strip(), len(heading.group(1)))
            index += 1
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            rows, next_index = _parse_markdown_table(lines, index)
            _add_word_table(document, rows)
            index = next_index
            continue

        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            _add_list_item(document, bullet.group(1), "List Bullet")
            index += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if numbered:
            _add_list_item(document, numbered.group(1), "List Number")
            index += 1
            continue

        _add_paragraph(document, line)
        index += 1


def _normalize_template_text(text: str) -> str:
    cleaned = _strip_markdown(text).replace("ё", "е").replace("Ё", "Е")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip(" .:-—–_").upper()


def _template_key_alias(value: str) -> str:
    normalized = re.sub(r"\s+", " ", value.strip().lower().replace("_", " "))
    normalized = normalized.replace("№", "N")
    if normalized in TEMPLATE_KEY_ALIASES:
        return TEMPLATE_KEY_ALIASES[normalized]
    compact = normalized.replace(" ", "").replace("-", "")
    if compact in {"pp", "пп", "ppalt", "ппalt"}:
        return "pp"
    if compact in {"di", "ди", "dialt", "диalt"}:
        return "di"
    if compact in {"ri", "ри", "rialt125", "риalt125"}:
        return "ri"
    return ""


def _official_template_key(
    meta: dict[str, str],
    content: str,
    title: str,
    document_type: str,
) -> str:
    explicit_template = _meta_value(meta, "template")
    if explicit_template:
        explicit_key = _template_key_alias(explicit_template)
        if explicit_key:
            return explicit_key

    haystack = f"{document_type}\n{title}\n{content[:1600]}".lower()
    compact = haystack.replace(" ", "").replace("-", "")

    if "должностная инструкция" in haystack or "диalt" in compact:
        return "di"
    if "положение о подраздел" in haystack or "ппalt" in compact:
        return "pp"
    if "рабочая инструкция" in haystack or "риalt125" in compact:
        return "ri"
    if document_type.strip().lower() == "положение":
        return "pp"
    return ""


def _official_approval_authority(meta: dict[str, str], template_key: str) -> str:
    authority = _meta_value(meta, "approval_authority").lower()
    if any(marker in authority for marker in ("совет", "директор", "board")):
        return "board"
    if any(marker in authority for marker in ("ректор", "председатель", "rector")):
        return "rector"

    explicit = _coerce_meta_bool(meta.get("approval_required"))
    if explicit is False:
        return "none"
    if template_key in {"pp", "di", "ri"}:
        return "rector"
    return "none"


def _delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _delete_body_range(document, start_marker: str, end_contains: str) -> None:
    paragraphs = list(document.paragraphs)
    start_index = None
    end_index = None

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if start_index is None and text == start_marker:
            start_index = index
        elif start_index is not None and end_contains in text:
            end_index = index
            break

    if start_index is None:
        return
    if end_index is None:
        end_index = start_index

    for paragraph in reversed(paragraphs[start_index : end_index + 1]):
        _delete_paragraph(paragraph)


def _is_template_hint_text(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text.strip().lower())
    if not normalized:
        return False
    return (
        "после заполнения прочерков" in normalized
        or "указывается в случае бумажного варианта" in normalized
        or "данная «шапка» оформляется" in normalized
        or "данная \"шапка\" оформляется" in normalized
        or normalized.startswith("*наименование должностей")
        or (
            normalized.startswith("(")
            and any(
                marker in normalized
                for marker in (
                    "полное наименование",
                    "сокращенное наименование",
                    "полное название",
                    "указать прочие",
                    "указывается цель",
                    "наименование подразделений",
                    "перечислить",
                    "описать специфическую",
                )
            )
        )
    )


def _remove_official_template_hints(document) -> None:
    for paragraph in list(document.paragraphs):
        if _is_template_hint_text(paragraph.text):
            _delete_paragraph(paragraph)


def _apply_official_approval_choice(document, template_key: str, meta: dict[str, str]) -> None:
    authority = _official_approval_authority(meta, template_key)
    if authority != "board":
        _delete_body_range(document, "УТВЕРЖДЕНО", "Данная «шапка» оформляется")
    if authority != "rector":
        _delete_body_range(document, "УТВЕРЖДАЮ", "Данная «шапка» оформляется на остальные")


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _iter_part_paragraphs(part):
    for paragraph in part.paragraphs:
        yield paragraph
    for table in part.tables:
        yield from _iter_table_paragraphs(table)


def _iter_all_template_paragraphs(document):
    yield from _iter_part_paragraphs(document)
    for section in document.sections:
        yield from _iter_part_paragraphs(section.header)
        yield from _iter_part_paragraphs(section.footer)


def _clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _copy_basic_run_format(source_run, target_run) -> None:
    if source_run is None:
        _set_run_font(target_run, size=14)
        return

    source_font = source_run.font
    target_font = target_run.font
    if source_font.name:
        target_font.name = source_font.name
        target_run._element.rPr.rFonts.set(qn("w:ascii"), source_font.name)
        target_run._element.rPr.rFonts.set(qn("w:hAnsi"), source_font.name)
        target_run._element.rPr.rFonts.set(qn("w:eastAsia"), source_font.name)
        target_run._element.rPr.rFonts.set(qn("w:cs"), source_font.name)
    else:
        _set_run_font(target_run, size=14)
    if source_font.size:
        target_font.size = source_font.size
    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.underline = source_run.underline
    if source_font.color and source_font.color.rgb:
        target_font.color.rgb = source_font.color.rgb


def _write_paragraph_text(paragraph, text: str) -> None:
    source_run = paragraph.runs[0] if paragraph.runs else None
    _clear_paragraph_content(paragraph)
    if text:
        run = paragraph.add_run(text)
        _copy_basic_run_format(source_run, run)


def _replace_blank_sequences(text: str, values: list[str]) -> str:
    value_iter = iter([value for value in values if value and value.lower() not in EMPTY_META_VALUES])

    def replacement(match):
        return next(value_iter, match.group(0))

    return re.sub(r"_{3,}", replacement, text)


def _template_year(meta: dict[str, str]) -> str:
    date_text = _meta_value(meta, "date")
    year_match = re.search(r"20\d{2}", date_text)
    if year_match:
        return year_match.group(0)
    return str(date.today().year)


def _official_template_values(meta: dict[str, str], template_key: str, title: str) -> dict[str, str]:
    department_full = _meta_value(meta, "department_full_name")
    department_short = _meta_value(meta, "department_short_name", default=department_full)
    position_full = _meta_value(meta, "position_full_name")
    position_short = _meta_value(meta, "position_short_name", default=position_full)
    title_value = _meta_value(meta, "title", default=title if title not in GENERIC_TITLES else "")

    if template_key == "pp" and not title_value and department_full:
        title_value = f"Положение о {department_full}"
    if template_key == "di" and not title_value:
        title_parts = [position_full, department_full]
        title_value = ", ".join(part for part in title_parts if part)
    if template_key == "ri" and not title_value:
        title_value = _meta_value(meta, "document_type", default="Рабочая инструкция")

    return {
        "department_full": department_full,
        "department_short": department_short,
        "position_full": position_full,
        "position_short": position_short,
        "title": title_value,
        "document_code": _meta_value(meta, "document_code"),
        "approver_name": _meta_value(meta, "approver_name", default="Ф.И.О."),
        "approver_position": _meta_value(meta, "approver_position", default="Председатель Правления-Ректор"),
        "date": _meta_value(meta, "date", default=f"«__» __________ 20__ г."),
        "board_protocol": _meta_value(meta, "board_protocol", default="Протокол №__ от «__»_________20__г."),
        "city_year": _meta_value(meta, "city_year", default=f"г. Алматы, {_template_year(meta)} год"),
        "supervisor_position": _meta_value(meta, "supervisor_position"),
        "leader_position": _meta_value(meta, "leader_position", "supervisor_position"),
        "extra_documents": _meta_value(meta, "extra_documents"),
        "purpose": _meta_value(meta, "purpose"),
    }


def _replace_official_template_text(text: str, template_key: str, values: dict[str, str]) -> str:
    original = text

    document_code = values["document_code"]
    if document_code:
        text = re.sub(r"\bПП\s*-\s*ALT\s*-\s*_{1,}", document_code, text)
        text = re.sub(r"\bПП-\s*ALT\s*-\s*_{1,}", document_code, text)
        text = re.sub(r"\bДИ\s*-\s*ALT\s*-?\s*_{0,}", document_code, text)
        text = re.sub(r"\bДИ-\s*ALT\s*-\s*_{1,}", document_code, text)
        text = re.sub(r"\bРИ\s*-\s*ALT\s*-\s*\d+", document_code, text)
        text = re.sub(r"\bРИ-\s*ALT\s*-\s*\d+", document_code, text)

    text = re.sub(
        r"____________\s*Ф\.И\.О\.",
        f"____________ {values['approver_name']}",
        text,
    )
    text = text.replace("Председатель Правления-Ректор", values["approver_position"])
    text = re.sub(r"«__»\s*__________\s*20__\s*г\.", values["date"], text)
    text = re.sub(r"Протокол №__ от «__»_________20__г\.", values["board_protocol"], text)
    text = re.sub(r"г\.\s*Алматы,\s*20\d{2}\s*год", values["city_year"], text)

    department_full = values["department_full"]
    department_short = values["department_short"]
    position_full = values["position_full"]
    position_short = values["position_short"]

    if template_key == "pp":
        if "Положение о" in text and "________________" in text:
            text = re.sub(r"Положение о\s*_{3,}", f"Положение о {department_full}", text)
        if "О (ОБ)" in text and "____" in text:
            text = f"О (ОБ) {department_full or '_______________________________'}"
        if "Настоящее Положение об" in text:
            text = _replace_blank_sequences(text, [department_full])
        if "определяет структуру" in text:
            text = _replace_blank_sequences(text, [department_full, department_short])
        if "является структурным подразделением" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "подчиняется непосредственно" in text:
            text = _replace_blank_sequences(text, [values["supervisor_position"]])
        if "Распределение обязанностей" in text or "осуществляет руководитель" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "ликвидация или его реорганизация" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "в своей работе руководствуется" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "Целью деятельности" in text:
            text = _replace_blank_sequences(text, [department_short, values["purpose"]])
        if "возглавляет" in text:
            text = _replace_blank_sequences(text, [department_short, values["leader_position"]])
        if "формируется согласно" in text or "работников" in text and "определяются" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "Основными задачами" in text or "осуществляет следующие функции" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "Работники" in text and "несут ответственность" in text:
            text = _replace_blank_sequences(text, [department_short])
        if "Степень ответственности" in text:
            text = _replace_blank_sequences(text, [department_short])

    if template_key == "di":
        title_value = values["title"] or ", ".join(part for part in (position_full, department_full) if part)
        if "Должностная инструкция" in text and "____" in text:
            text = re.sub(r"Должностная инструкция\s*_{3,}", f"Должностная инструкция {title_value}", text)
        if text.strip().startswith("___") and title_value:
            text = title_value
        if "Настоящая Должностная инструкция" in text:
            text = _replace_blank_sequences(text, [position_full])
        if "Основными документами" in text:
            text = _replace_blank_sequences(text, [position_full, position_short])
        if text.strip().startswith("об") and "____" in text:
            text = _replace_blank_sequences(text, [department_full, department_short])
        if "подчиняется непосредственно" in text:
            text = _replace_blank_sequences(text, [position_short, values["supervisor_position"]])
        if "назначается на должность" in text:
            text = _replace_blank_sequences(text, [position_short])
        if "по представлению" in text:
            text = _replace_blank_sequences(text, [values["supervisor_position"]])
        if "В своей деятельности" in text or "Во время отсутствия" in text:
            text = _replace_blank_sequences(text, [position_short])
        if "Положением об" in text:
            text = _replace_blank_sequences(text, [department_full])
        if "На должность" in text:
            text = _replace_blank_sequences(text, [position_short])
        if "имеет право" in text or "несет ответственность" in text:
            text = _replace_blank_sequences(text, [position_short])

    if template_key == "ri" and values["title"]:
        if original.strip().upper() in {"ПОРЯДОК", "РАБОЧАЯ ИНСТРУКЦИЯ"}:
            return text
        if "ПОРЯДОК" in text and "РАЗРАБОТКИ" not in text:
            text = values["title"]

    text = re.sub(
        r"\s*\([^)]*(?:полное|сокращенное)\s+(?:наименование|название)[^)]*\)",
        "",
        text,
        flags=re.IGNORECASE,
    )

    return text


def _apply_official_template_replacements(
    document,
    template_key: str,
    meta: dict[str, str],
    title: str,
) -> None:
    values = _official_template_values(meta, template_key, title)
    for paragraph in _iter_all_template_paragraphs(document):
        original = paragraph.text
        if not original.strip():
            continue
        updated = _replace_official_template_text(original, template_key, values)
        if updated != original:
            _write_paragraph_text(paragraph, updated)


def _canonical_template_heading(text: str, template_key: str) -> str:
    normalized = _normalize_template_text(text)
    for heading in TEMPLATE_SECTION_HEADINGS.get(template_key, ()):
        if normalized == _normalize_template_text(heading):
            return heading
    for heading in TEMPLATE_STOP_HEADINGS:
        if normalized == _normalize_template_text(heading):
            return heading
    return ""


def _parse_markdown_template_sections(content: str, template_key: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    current_heading = ""
    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        heading_match = re.match(r"^\s{0,3}#{1,6}\s+(.+)$", line)
        if heading_match:
            canonical = _canonical_template_heading(heading_match.group(1), template_key)
            if canonical and canonical not in TEMPLATE_STOP_HEADINGS:
                current_heading = canonical
                sections.setdefault(current_heading, [])
                continue
        if current_heading:
            sections[current_heading].append(line)

    return {
        heading: _trim_blank_lines(lines)
        for heading, lines in sections.items()
        if _trim_blank_lines(lines)
    }


def _trim_blank_lines(lines: list[str]) -> list[str]:
    start = 0
    end = len(lines)
    while start < end and not lines[start].strip():
        start += 1
    while end > start and not lines[end - 1].strip():
        end -= 1
    return lines[start:end]


def _insert_paragraph_after(paragraph, text: str = ""):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    if text:
        run = new_paragraph.add_run(text)
        _set_run_font(run, size=14)
    return new_paragraph


def _insert_markdown_lines_after(anchor, lines: list[str]) -> None:
    cursor = anchor
    for line in lines:
        stripped = line.strip()
        if not stripped:
            cursor = _insert_paragraph_after(cursor)
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            cursor = _insert_paragraph_after(cursor, _strip_markdown(heading.group(2)))
            cursor.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in cursor.runs:
                _set_run_font(run, bold=True, size=14)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        text = bullet.group(1) if bullet else numbered.group(1) if numbered else stripped
        cursor = _insert_paragraph_after(cursor, _strip_markdown(text))
        cursor.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        try:
            if bullet:
                cursor.style = "List Bullet"
            elif numbered:
                cursor.style = "List Number"
        except KeyError:
            pass


def _replace_official_template_sections(document, template_key: str, content: str) -> None:
    sections = _parse_markdown_template_sections(content, template_key)
    if not sections:
        return

    for heading, lines in sections.items():
        paragraphs = list(document.paragraphs)
        heading_index = None
        for index, paragraph in enumerate(paragraphs):
            if _canonical_template_heading(paragraph.text, template_key) == heading:
                heading_index = index
                break
        if heading_index is None:
            continue

        end_index = len(paragraphs)
        for index in range(heading_index + 1, len(paragraphs)):
            candidate = _canonical_template_heading(paragraphs[index].text, template_key)
            if candidate in TEMPLATE_SECTION_HEADINGS.get(template_key, ()) or candidate in TEMPLATE_STOP_HEADINGS:
                end_index = index
                break

        anchor = paragraphs[heading_index]
        for paragraph in reversed(paragraphs[heading_index + 1 : end_index]):
            _delete_paragraph(paragraph)
        _insert_markdown_lines_after(anchor, lines)


def _build_official_template_docx(
    content: str,
    title: str,
    meta: dict[str, str],
    template_key: str,
) -> bytes:
    template_path = OFFICIAL_DOCX_TEMPLATES.get(template_key)
    if not template_path or not template_path.exists():
        return _build_formal_docx(content, title, meta)

    document = DocxDocument(str(template_path))
    _apply_official_approval_choice(document, template_key, meta)
    _remove_official_template_hints(document)
    _apply_official_template_replacements(document, template_key, meta, title)
    _replace_official_template_sections(document, template_key, content)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_formal_docx(content: str, title: str, meta: dict[str, str]) -> bytes:
    document = DocxDocument()
    _configure_docx_styles(document)
    _apply_formal_page_setup(document)

    document_type = _document_type_from(title, content, meta)
    document_title = _document_title_from(title, content, meta, document_type)
    requires_approval = _approval_required(meta, document_type, document_title, content)
    standalone_title_page = requires_approval and not _is_letter_like_document(document_type, document_title)

    if requires_approval:
        _add_approval_block(document, meta)
    else:
        _add_address_block(document, meta)

    _add_formal_title_block(
        document,
        document_type,
        document_title,
        standalone_title_page=standalone_title_page,
    )

    body_content = _remove_leading_duplicate_title(content, document_type, document_title)
    if body_content:
        if standalone_title_page:
            document.add_page_break()
        else:
            _add_spacer(document, 18)
        _append_markdown_to_docx(document, body_content)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _configure_docx_styles(document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.bold = True
        style.font.size = Pt(14)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)


def _new_docx_document():
    if DOCX_TEMPLATE_PATH.exists():
        document = DocxDocument(str(DOCX_TEMPLATE_PATH))
    else:
        document = DocxDocument()
    _configure_docx_styles(document)
    return document


def _normalized_heading(text: str) -> str:
    return re.sub(r"\s+", " ", _strip_markdown(text)).strip().lower()


def _content_already_starts_with_title(content: str, title: str) -> bool:
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,6}\s+(.+)$", line)
        return bool(heading and _normalized_heading(heading.group(1)) == _normalized_heading(title))
    return False


def _add_title_if_needed(document, content: str, title: str) -> None:
    cleaned_title = title.strip()
    if not cleaned_title or cleaned_title in GENERIC_TITLES:
        return
    if _content_already_starts_with_title(content, cleaned_title):
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(_strip_markdown(cleaned_title))
    _set_run_font(run, bold=True, size=14)


def _add_heading(document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(_strip_markdown(text))
    _set_run_font(run, bold=True, size=14)


def _add_paragraph(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(_strip_markdown(text))
    _set_run_font(run, size=14)


def _add_list_item(document, text: str, style: str) -> None:
    paragraph = document.add_paragraph()
    if style in document.styles:
        paragraph.style = style
    run = paragraph.add_run(_strip_markdown(text))
    _set_run_font(run, size=14)


def _is_markdown_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def _parse_markdown_table(lines: list[str], start_index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start_index
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            break
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if not _is_markdown_table_separator(cells):
            rows.append(cells)
        index += 1
    return rows, index


def _add_word_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    if "Table Grid" in document.styles:
        table.style = "Table Grid"

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for column_index in range(column_count):
            cell_text = row_values[column_index] if column_index < len(row_values) else ""
            paragraph = row.cells[column_index].paragraphs[0]
            run = paragraph.add_run(_strip_markdown(cell_text))
            _set_run_font(run, bold=row_index == 0, size=12)

    document.add_paragraph()


def _build_docx(content: str, title: str) -> bytes:
    meta, clean_content = _extract_docx_meta(content)
    document_type = _document_type_from(title, clean_content, meta)
    official_template_key = _official_template_key(meta, clean_content, title, document_type)
    if official_template_key:
        return _build_official_template_docx(
            clean_content,
            title,
            meta,
            official_template_key,
        )

    if _should_use_formal_docx(meta, clean_content, title):
        return _build_formal_docx(clean_content, title, meta)

    document = _new_docx_document()
    document.add_paragraph()
    _add_title_if_needed(document, clean_content, title)
    _append_markdown_to_docx(document, clean_content)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _markdown_table_rows(content: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            if rows:
                break
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def _build_xlsx(content: str, title: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Answer"

    rows = _markdown_table_rows(content)
    if rows:
        for row in rows:
            sheet.append(row)
    else:
        sheet.append([title])
        sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([])
        for line in _strip_markdown(content).splitlines():
            sheet.append([line])

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column_cells in sheet.columns:
        values = [str(cell.value or "") for cell in column_cells]
        max_len = min(max((len(value) for value in values), default=12), 80)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(12, max_len + 2)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_csv(content: str) -> bytes:
    output = StringIO()
    writer = csv.writer(output)
    rows = _markdown_table_rows(content)
    if rows:
        writer.writerows(rows)
    else:
        for line in _strip_markdown(content).splitlines():
            writer.writerow([line])
    return output.getvalue().encode("utf-8-sig")


@router.post("/{file_format}")
async def export_answer(
    file_format: ExportFormat,
    request: ExportRequest,
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """
    Export an assistant answer to a downloadable office/text file.
    """
    if not request.content.strip():
        raise HTTPException(status_code=400, detail="Export content is empty")

    title = request.title.strip() or "ALTBot answer"

    if file_format == "docx":
        filename = _safe_filename(title, "docx")
        return _download_response(
            _build_docx(request.content, title),
            filename,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    if file_format == "xlsx":
        filename = _safe_filename(title, "xlsx")
        return _download_response(
            _build_xlsx(request.content, title),
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    if file_format == "txt":
        filename = _safe_filename(title, "txt")
        return _download_response(
            _strip_markdown(request.content).encode("utf-8"),
            filename,
            "text/plain; charset=utf-8",
        )

    if file_format == "md":
        filename = _safe_filename(title, "md")
        return _download_response(
            request.content.encode("utf-8"),
            filename,
            "text/markdown; charset=utf-8",
        )

    if file_format == "csv":
        filename = _safe_filename(title, "csv")
        return _download_response(
            _build_csv(request.content),
            filename,
            "text/csv; charset=utf-8",
        )

    raise HTTPException(status_code=400, detail="Unsupported export format")
